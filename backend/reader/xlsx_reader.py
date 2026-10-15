from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET

from .contract import ReaderInput, ReaderResult, lineage, new_result
from .utils import safe_text, sha256_bytes

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


XLSX_READER_VERSION = "xlsx-ooxml/0.2.0"
XLSX_PARSER_VERSION = "openpyxl+ooxml+Pillow/2"
XLSX_FALLBACK_PARSER_VERSION = "stdlib-ooxml+Pillow/1"

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _local(element: ET.Element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _child(element: ET.Element, name: str) -> ET.Element | None:
    return next((x for x in list(element) if _local(x) == name), None)


def _children(element: ET.Element, name: str) -> list[ET.Element]:
    return [x for x in list(element) if _local(x) == name]


def _target(base: str, target: str) -> str:
    return posixpath.normpath(posixpath.join(posixpath.dirname(base), target.lstrip("/"))) if not target.startswith("/") else target.lstrip("/")


def _xml(z: zipfile.ZipFile, name: str) -> ET.Element:
    return ET.fromstring(z.read(name))


def _rels(z: zipfile.ZipFile, source: str) -> dict[str, dict[str, str]]:
    rel_path = posixpath.join(posixpath.dirname(source), "_rels", posixpath.basename(source) + ".rels")
    if rel_path not in z.namelist():
        return {}
    root = _xml(z, rel_path)
    return {x.attrib.get("Id", ""): {"type": x.attrib.get("Type", ""), "target": _target(source, x.attrib.get("Target", "")), "mode": x.attrib.get("TargetMode", "")} for x in root if _local(x) == "Relationship"}


def _xml_text(element: ET.Element | None) -> str:
    return "".join(element.itertext()).strip() if element is not None else ""


def _column_number(coordinate: str) -> int:
    letters = re.match(r"[A-Z]+", coordinate.upper())
    if not letters:
        return 0
    value = 0
    for char in letters.group(0):
        value = value * 26 + ord(char) - 64
    return value


def _number(value: str | None) -> Any:
    if value is None or value == "":
        return None
    try:
        number = float(value)
        return int(number) if number.is_integer() else number
    except ValueError:
        return value


def _shared_strings(z: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    return [_xml_text(x) for x in _xml(z, "xl/sharedStrings.xml") if _local(x) == "si"]


def _style_records(z: zipfile.ZipFile) -> tuple[list[dict[str, Any]], set[int]]:
    if "xl/styles.xml" not in z.namelist():
        return [], set()
    root = _xml(z, "xl/styles.xml")
    fonts = _child(root, "fonts")
    fills = _child(root, "fills")
    numfmts_node = _child(root, "numFmts")
    numfmts = {int(x.attrib.get("numFmtId", "0")): x.attrib.get("formatCode", "General") for x in _children(numfmts_node if numfmts_node is not None else ET.Element("numFmts"), "numFmt")}
    builtins = {9: "0%", 10: "0.00%", 14: "mm-dd-yy"}
    xfs = _child(root, "cellXfs")
    records: list[dict[str, Any]] = []
    for xf in _children(xfs if xfs is not None else ET.Element("cellXfs"), "xf"):
        font_id = int(xf.attrib.get("fontId", "0") or 0)
        fill_id = int(xf.attrib.get("fillId", "0") or 0)
        num_id = int(xf.attrib.get("numFmtId", "0") or 0)
        font = _children(fonts if fonts is not None else ET.Element("fonts"), "font")[font_id] if fonts is not None and font_id < len(_children(fonts, "font")) else None
        fill = _children(fills if fills is not None else ET.Element("fills"), "fill")[fill_id] if fills is not None and fill_id < len(_children(fills, "fill")) else None
        color = _child(font, "color") if font is not None else None
        pattern = _child(fill, "patternFill") if fill is not None else None
        fg = _child(pattern, "fgColor") if pattern is not None else None
        bg = _child(pattern, "bgColor") if pattern is not None else None
        alignment = _child(xf, "alignment")
        protection = _child(xf, "protection")
        records.append({
            "number_format": numfmts.get(num_id, builtins.get(num_id, "General")),
            "style": {"font": {"name": (_child(font, "name").attrib.get("val") if font is not None and _child(font, "name") is not None else None), "size": _number(_child(font, "sz").attrib.get("val") if font is not None and _child(font, "sz") is not None else None), "bold": font is not None and _child(font, "b") is not None, "italic": font is not None and _child(font, "i") is not None, "color": (color.attrib.get("rgb", "") if color is not None else "")}, "fill": {"fill_type": (pattern.attrib.get("patternType") if pattern is not None else None), "fgColor": (fg.attrib.get("rgb", "") if fg is not None else ""), "bgColor": (bg.attrib.get("rgb", "") if bg is not None else "")}, "alignment": {"horizontal": alignment.attrib.get("horizontal") if alignment is not None else None, "vertical": alignment.attrib.get("vertical") if alignment is not None else None, "wrap_text": alignment.attrib.get("wrapText") == "1" if alignment is not None and "wrapText" in alignment.attrib else None}},
            "protection": {"locked": protection.attrib.get("locked") != "0" if protection is not None else True},
        })
    return records, set(range(len(records)))


def _chart_raw(z: zipfile.ZipFile, path: str, index: int) -> dict[str, Any]:
    root = _xml(z, path)
    plot = next((x for x in root.iter() if _local(x) == "plotArea"), None)
    chart = next((x for x in list(plot if plot is not None else ET.Element("plotArea")) if _local(x).endswith("Chart")), None)
    series: list[dict[str, Any]] = []
    for i, ser in enumerate([x for x in (list(chart) if chart is not None else []) if _local(x) == "ser"]):
        def ref(name: str) -> str:
            node = next((x for x in ser.iter() if _local(x) == name), None)
            formula = next((x for x in (list(node if node is not None else ET.Element("ref"))) if _local(x) == "f"), None)
            return _xml_text(formula)
        title = next((x for x in ser.iter() if _local(x) == "v"), None)
        series.append({"index": i, "title": _xml_text(title), "categories_formula": ref("cat"), "values_formula": ref("val"), "x_values_formula": ref("xVal"), "y_values_formula": ref("yVal")})
    return {"index": index, "chart_type": _local(chart) if chart is not None else "unknown", "series": series, "package_path": path, "sha256": sha256_bytes(z.read(path)), "raw_xml_preserved": True}


def _drawing_objects(z: zipfile.ZipFile, drawing_path: str, rels: dict[str, dict[str, str]]) -> tuple[list[tuple[str, dict[str, Any]]], list[tuple[str, dict[str, Any]]]]:
    root = _xml(z, drawing_path)
    charts: list[tuple[str, dict[str, Any]]] = []
    images: list[tuple[str, dict[str, Any]]] = []
    for anchor in [x for x in root if _local(x) in {"twoCellAnchor", "oneCellAnchor", "absoluteAnchor"}]:
        graphic = next((x for x in anchor.iter() if _local(x) == "graphicFrame"), None)
        if graphic is not None:
            chart = next((x for x in graphic.iter() if _local(x) == "chart"), None)
            if chart is not None:
                rid = chart.attrib.get(f"{{{_DOC_REL_NS}}}id", "")
                if rid in rels:
                    charts.append((rels[rid]["target"], {"anchor": _drawing_anchor(anchor), "relationship_id": rid}))
        pic = next((x for x in anchor.iter() if _local(x) == "pic"), None)
        if pic is not None:
            blip = next((x for x in pic.iter() if _local(x) == "blip"), None)
            rid = (blip.attrib.get(f"{{{_DOC_REL_NS}}}embed", "") if blip is not None else "")
            if rid in rels:
                path = rels[rid]["target"]
                images.append((path, {"anchor": _drawing_anchor(anchor), "relationship_id": rid}))
    return charts, images


def _drawing_anchor(anchor: ET.Element) -> dict[str, Any]:
    out: dict[str, Any] = {"type": _local(anchor)}
    for label in ("from", "to", "ext"):
        node = next((x for x in anchor if _local(x) == label), None)
        if node is not None:
            out[label] = {(_local(x)): _number(x.text) for x in node}
    return out


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _display_value(value: Any, number_format: str) -> Any:
    # Do not transform raw values. This is only a parser-derived display hint.
    if value is None:
        return None
    if isinstance(value, float) and number_format and "%" in number_format:
        return f"{value * 100:g}%"
    return _json_value(value)


def _anchor(anchor: Any) -> dict[str, Any]:
    out: dict[str, Any] = {"type": type(anchor).__name__}
    for label in ("_from", "_to"):
        point = getattr(anchor, label, None)
        if point is not None:
            out[label[1:]] = {"col": getattr(point, "col", None), "row": getattr(point, "row", None), "colOff": getattr(point, "colOff", None), "rowOff": getattr(point, "rowOff", None)}
    ext = getattr(anchor, "ext", None)
    if ext is not None:
        out["ext"] = {"cx": getattr(ext, "cx", None), "cy": getattr(ext, "cy", None)}
    return out


def _series_ref(series: Any, attr: str) -> str:
    try:
        ref = getattr(series, attr)
        ref = getattr(ref, "numRef", None) or getattr(ref, "strRef", None)
        return str(getattr(ref, "f", "") or "")
    except Exception:
        return ""


def _chart_record(chart: Any, index: int) -> dict[str, Any]:
    out: dict[str, Any] = {"index": index, "chart_type": type(chart).__name__, "anchor": _anchor(chart.anchor), "series": []}
    for i, series in enumerate(getattr(chart, "ser", []) or []):
        out["series"].append({"index": i, "title": str(getattr(series, "tx", "") or ""), "categories_formula": _series_ref(series, "cat"), "values_formula": _series_ref(series, "val"), "x_values_formula": _series_ref(series, "xVal"), "y_values_formula": _series_ref(series, "yVal")})
    return out


def _sheet_png(result: ReaderResult, sheet: dict[str, Any], out: Path) -> None:
    from PIL import Image, ImageDraw
    max_rows, max_cols = min(max(int(sheet.get("max_row", 1)), 1), 100), min(max(int(sheet.get("max_column", 1)), 1), 30)
    cell_w, cell_h = 150, 28
    image = Image.new("RGB", (min(1800, (max_cols + 1) * cell_w), min(3000, (max_rows + 2) * cell_h)), "white")
    draw = ImageDraw.Draw(image)
    draw.text((8, 5), f"Sheet: {sheet['title']}  hidden={sheet['hidden']}", fill="black")
    for c in range(1, max_cols + 1):
        x = c * cell_w
        draw.rectangle((x, cell_h, x + cell_w, cell_h * 2), outline="#888", fill="#e8eef8")
        draw.text((x + 4, cell_h + 6), str(c), fill="black")
    by_coord = {(int(x["row"]), int(x["column"])): x for x in sheet["cells"]}
    for r in range(1, max_rows + 1):
        y = (r + 1) * cell_h
        draw.text((5, y + 6), str(r), fill="black")
        for c in range(1, max_cols + 1):
            x = c * cell_w
            cell = by_coord.get((r, c), {})
            fill = "#fff8dc" if cell.get("formula") else "white"
            draw.rectangle((x, y, x + cell_w, y + cell_h), outline="#c0c0c0", fill=fill)
            value = cell.get("display_value")
            if value is not None:
                draw.text((x + 4, y + 6), safe_text(value)[:22], fill="black")
    out.parent.mkdir(parents=True, exist_ok=True); image.save(out, "PNG")
    result.visual_pages.append({"kind": "sheet", "sheet": sheet["title"], "path": str(out), "sha256": sha256_bytes(out.read_bytes()), "dimensions": {"width": image.width, "height": image.height}, "render_mode": "pillow-structural-proxy", "bounded_inventory": {"max_rows": max_rows, "max_columns": max_cols}})
    result.derived_artifacts.append(lineage(result.raw_artifact_sha256, result.source_version, out, "visual-sheet", render_mode="pillow-structural-proxy"))


def _read_xlsx_ooxml(inp: ReaderInput, output_dir: str | Path, *, parse_warning: str = "") -> ReaderResult:
    """Parse OOXML directly when a library rejects a valid producer variant.

    This route reads the committed package parts without rewriting the source
    or normalizing unsupported chart enum values.  It therefore retains raw
    formula/style/chart/image evidence while remaining deterministic.
    """
    result = new_result(inp, "XLSX")
    result.reader_version = XLSX_READER_VERSION
    result.parser_version = XLSX_FALLBACK_PARSER_VERSION
    if parse_warning:
        result.warnings.append(f"openpyxl route unavailable; OOXML fallback used: {parse_warning}")
    try:
        with zipfile.ZipFile(inp.path) as z:
            names = set(z.namelist())
            workbook_path = "xl/workbook.xml"
            workbook = _xml(z, workbook_path)
            workbook_rels = _rels(z, workbook_path)
            shared = _shared_strings(z)
            styles, valid_style_ids = _style_records(z)
            sheets: list[dict[str, Any]] = []
            all_charts: list[dict[str, Any]] = []
            all_images: list[dict[str, Any]] = []
            for index, sheet_node in enumerate([x for x in workbook.iter() if _local(x) == "sheet"], 1):
                rid = sheet_node.attrib.get(f"{{{_DOC_REL_NS}}}id", "")
                relation = workbook_rels.get(rid, {})
                sheet_path = relation.get("target", "")
                if sheet_path not in names:
                    raise ValueError(f"worksheet target is missing: {sheet_path}")
                root = _xml(z, sheet_path)
                dimension = _child(root, "dimension")
                dimension_ref = dimension.attrib.get("ref", "") if dimension is not None else ""
                max_row, max_col = 1, 1
                if ":" in dimension_ref:
                    end = dimension_ref.rsplit(":", 1)[-1]
                    max_row = int(re.search(r"\d+", end).group(0)) if re.search(r"\d+", end) else 1
                    max_col = _column_number(end)
                elif dimension_ref:
                    max_row = int(re.search(r"\d+", dimension_ref).group(0)) if re.search(r"\d+", dimension_ref) else 1
                    max_col = _column_number(dimension_ref)
                sheet = {"order": index, "title": sheet_node.attrib.get("name", f"Sheet{index}"), "hidden": sheet_node.attrib.get("state", "visible") != "visible", "sheet_state": sheet_node.attrib.get("state", "visible"), "max_row": max_row, "max_column": max_col, "merged_ranges": [], "row_dimensions": {}, "column_dimensions": {}, "tables": [], "conditional_formatting": [], "cells": [], "charts": [], "images": [], "hyperlinks": [], "comments": []}
                merge_cells = _child(root, "mergeCells")
                sheet["merged_ranges"] = [x.attrib.get("ref", "") for x in _children(merge_cells if merge_cells is not None else ET.Element("mergeCells"), "mergeCell")]
                cols = _child(root, "cols")
                for col in _children(cols if cols is not None else ET.Element("cols"), "col"):
                    first, last = col.attrib.get("min", ""), col.attrib.get("max", "")
                    key = first if first == last else f"{first}:{last}"
                    sheet["column_dimensions"][key] = {"hidden": col.attrib.get("hidden") == "1", "width": _number(col.attrib.get("width")), "outline_level": int(col.attrib.get("outlineLevel", "0") or 0)}
                sheet_data = _child(root, "sheetData")
                for row in _children(sheet_data if sheet_data is not None else ET.Element("sheetData"), "row"):
                    row_num = int(row.attrib.get("r", "0") or 0)
                    if row.attrib.get("ht") is not None or row.attrib.get("hidden") == "1" or row.attrib.get("outlineLevel"):
                        sheet["row_dimensions"][str(row_num)] = {"hidden": row.attrib.get("hidden") == "1", "height": _number(row.attrib.get("ht")), "outline_level": int(row.attrib.get("outlineLevel", "0") or 0)}
                    for cell in _children(row, "c"):
                        coordinate = cell.attrib.get("r", "")
                        col_num = _column_number(coordinate)
                        row_num_cell = int(re.search(r"\d+", coordinate).group(0)) if re.search(r"\d+", coordinate) else row_num
                        style_id = int(cell.attrib.get("s", "0") or 0)
                        cell_type = cell.attrib.get("t", "n")
                        value_node = _child(cell, "v")
                        formula_node = _child(cell, "f")
                        raw = _xml_text(value_node)
                        if cell_type == "s":
                            raw_value: Any = shared[int(raw)] if raw.isdigit() and int(raw) < len(shared) else raw
                        elif cell_type == "inlineStr":
                            raw_value = _xml_text(_child(cell, "is"))
                        elif cell_type == "b":
                            raw_value = raw == "1"
                        elif cell_type in {"str", "e"}:
                            raw_value = raw
                        else:
                            raw_value = _number(raw)
                        formula = f"={_xml_text(formula_node)}" if formula_node is not None else None
                        if formula is not None:
                            raw_value = formula
                        style_record = styles[style_id] if style_id < len(styles) else {"number_format": "General", "style": {}, "protection": {"locked": True}}
                        item = {"coordinate": coordinate, "row": row_num_cell, "column": col_num, "raw_value": raw_value, "display_value": _display_value(_number(raw) if formula is not None and cell_type == "n" else raw_value, style_record.get("number_format", "General")), "formula": formula, "data_type": "f" if formula is not None else cell_type, "number_format": style_record.get("number_format", "General"), "style_id": style_id, "style": style_record.get("style", {}), "protection": style_record.get("protection", {"locked": True})}
                        if coordinate:
                            sheet["cells"].append(item)
                        if style_id not in valid_style_ids and styles:
                            result.warnings.append(f"XLSX cell {sheet['title']}!{coordinate} references unknown style id {style_id}")
                conditional = _child(root, "conditionalFormatting")
                for cf in _children(conditional if conditional is not None else ET.Element("conditionalFormatting"), "cfRule"):
                    sheet["conditional_formatting"].append({"range": (_child(root, "conditionalFormatting").attrib.get("sqref", "") if _child(root, "conditionalFormatting") is not None else ""), "rules": [cf.attrib.get("type", "cfRule")]})
                sheet_rels = _rels(z, sheet_path)
                for rel in sheet_rels.values():
                    if rel.get("type", "").endswith("/table") and rel["target"] in names:
                        table_root = _xml(z, rel["target"])
                        style = _child(table_root, "tableStyleInfo")
                        sheet["tables"].append({"name": table_root.attrib.get("name", ""), "ref": table_root.attrib.get("ref", ""), "display_name": table_root.attrib.get("displayName", ""), "style": style.attrib.get("name", "") if style is not None else ""})
                    if rel.get("type", "").endswith("/drawing") and rel["target"] in names:
                        drawing_rels = _rels(z, rel["target"])
                        charts, images = _drawing_objects(z, rel["target"], drawing_rels)
                        for chart_path, evidence in charts:
                            record = _chart_raw(z, chart_path, len(sheet["charts"]))
                            record["sheet"] = sheet["title"]
                            record.update(evidence)
                            sheet["charts"].append(record)
                            all_charts.append(record)
                            result.embedded_objects.append({"kind": "chart", "sheet": sheet["title"], "index": record["index"], "raw_data": record, "visual_required": True})
                        for image_path, evidence in images:
                            blob = z.read(image_path) if image_path in names else b""
                            info = {"kind": "image", "sheet": sheet["title"], "index": len(sheet["images"]), "sha256": sha256_bytes(blob) if blob else None, "size": len(blob), "anchor": evidence.get("anchor", {}), "package_path": image_path, "relationship_id": evidence.get("relationship_id", "")}
                            sheet["images"].append(info)
                            all_images.append(info)
                            result.embedded_objects.append(info)
                sheets.append(sheet)
            raw_drawings = [{"path": n, "sha256": sha256_bytes(z.read(n)), "size": z.getinfo(n).file_size} for n in names if n.startswith("xl/drawings/") and n.endswith(".xml")]
            chart_parts = [{"path": n, "sha256": sha256_bytes(z.read(n)), "size": z.getinfo(n).file_size} for n in names if n.startswith("xl/charts/") and n.endswith(".xml")]
        for sheet in sheets:
            _sheet_png(result, sheet, Path(output_dir) / "sheets" / f"sheet-{sheet['order']:04d}.png")
        formula_count = sum(1 for s in sheets for c in s["cells"] if c.get("formula"))
        styles_ok = bool(styles) and all(int(c.get("style_id", 0)) in valid_style_ids for s in sheets for c in s["cells"])
        charts_ok = len(all_charts) == len(chart_parts) and all(bool(x.get("series") or x.get("raw_xml_preserved")) for x in all_charts)
        images_ok = all(x.get("sha256") for x in all_images)
        result.structured_semantics = {"sheet_count": len(sheets), "sheets": sheets, "named_ranges": [], "drawing_parts": raw_drawings, "chart_parts": chart_parts, "ooxml_route": "stdlib-fallback"}
        result.fidelity_flags = {"raw_data_fidelity": True, "formula_fidelity": True, "display_values_parser_derived": True, "sheet_order_and_hidden_preserved": True, "merged_ranges_preserved": True, "dimensions_preserved": True, "styles_accounted": styles_ok, "charts_accounted": charts_ok, "chart_visual_evidence": all(any(v.get("sheet") == x.get("sheet") for v in result.visual_pages) for x in all_charts), "images_accounted": images_ok, "visual_evidence_available": bool(result.visual_pages), "formula_count": formula_count, "chart_count": len(all_charts), "image_count": len(all_images), "failure_explicit": True}
        result.warnings.append("sheet PNGs are deterministic structural visual proxies derived from OOXML cells; chart/image raw package evidence is retained")
    except Exception as exc:
        result.warnings.append(f"XLSX OOXML fallback parse failed closed: {exc}")
        result.fidelity_flags = {"raw_data_fidelity": False, "formula_fidelity": False, "styles_accounted": False, "charts_accounted": False, "images_accounted": False, "visual_evidence_available": False, "failure_explicit": True}
    return result.finalize()


def _read_xlsx_openpyxl(inp: ReaderInput, output_dir: str | Path) -> ReaderResult:
    result = new_result(inp, "XLSX")
    result.reader_version = XLSX_READER_VERSION
    result.parser_version = XLSX_PARSER_VERSION
    if load_workbook is None:
        return _read_xlsx_ooxml(inp, output_dir, parse_warning="openpyxl is unavailable")
    try:
        wb = load_workbook(inp.path, data_only=False, read_only=False)
        try:
            values_wb = load_workbook(inp.path, data_only=True, read_only=False)
        except Exception:
            values_wb = None
        sheets = []
        for index, ws in enumerate(wb.worksheets, 1):
            values_ws = values_wb[ws.title] if values_wb and ws.title in values_wb.sheetnames else None
            sheet = {"order": index, "title": ws.title, "hidden": ws.sheet_state != "visible", "sheet_state": ws.sheet_state, "max_row": ws.max_row, "max_column": ws.max_column, "merged_ranges": [str(x) for x in ws.merged_cells.ranges], "row_dimensions": {}, "column_dimensions": {}, "tables": [], "conditional_formatting": [], "cells": [], "charts": [], "images": [], "hyperlinks": [], "comments": []}
            for key, dim in ws.row_dimensions.items():
                if dim.hidden or dim.height is not None or dim.outlineLevel:
                    sheet["row_dimensions"][str(key)] = {"hidden": bool(dim.hidden), "height": dim.height, "outline_level": dim.outlineLevel}
            for key, dim in ws.column_dimensions.items():
                if dim.hidden or dim.width is not None or dim.outline_level:
                    sheet["column_dimensions"][str(key)] = {"hidden": bool(dim.hidden), "width": dim.width, "outline_level": dim.outline_level}
            # TableList.items() changed shape across openpyxl releases: in
            # some versions its values are refs rather than Table objects.
            # Indexing by name is the stable API and retains the real table.
            for name in ws.tables.keys():
                table = ws.tables[name]
                sheet["tables"].append({"name": name, "ref": table.ref, "display_name": table.displayName, "style": table.tableStyleInfo.name if table.tableStyleInfo else ""})
            for cf in ws.conditional_formatting:
                sheet["conditional_formatting"].append({"range": str(cf), "rules": [type(rule).__name__ for rule in ws.conditional_formatting[cf]]})
            # Iterate materialized cells, not the rectangular max_row*max_column
            # grid. Formatting-only sheets can advertise millions of empty
            # coordinates and must not exhaust a 2GB host.
            if int(ws.max_row or 0) * int(ws.max_column or 0) > 250000:
                result.warnings.append(f"XLSX sheet {ws.title} has a large advertised grid; sparse materialized-cell traversal used")
            for cell in sorted(getattr(ws, "_cells", {}).values(), key=lambda x: (x.row, x.column)):
                if cell.value is None and not cell.has_style and not cell.hyperlink and not cell.comment:
                    continue
                cached = values_ws[cell.coordinate].value if values_ws is not None else None
                item = {"coordinate": cell.coordinate, "row": cell.row, "column": cell.column, "raw_value": _json_value(cell.value), "display_value": _display_value(cached if cell.data_type == "f" else cell.value, cell.number_format), "formula": cell.value if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")) else None, "data_type": cell.data_type, "number_format": cell.number_format, "style_id": cell.style_id, "style": {"font": {"name": cell.font.name, "size": cell.font.sz, "bold": cell.font.bold, "italic": cell.font.italic, "color": str(cell.font.color.rgb) if cell.font.color and cell.font.color.type == "rgb" and cell.font.color.rgb else ""}, "fill": {"fill_type": cell.fill.fill_type, "fgColor": str(cell.fill.fgColor.rgb) if cell.fill.fgColor.type == "rgb" and cell.fill.fgColor.rgb else "", "bgColor": str(cell.fill.bgColor.rgb) if cell.fill.bgColor.type == "rgb" and cell.fill.bgColor.rgb else ""}, "alignment": {"horizontal": cell.alignment.horizontal, "vertical": cell.alignment.vertical, "wrap_text": cell.alignment.wrap_text}}, "protection": {"locked": cell.protection.locked}}
                if cell.hyperlink:
                    item["hyperlink"] = {"target": cell.hyperlink.target, "location": cell.hyperlink.location, "tooltip": cell.hyperlink.tooltip}; sheet["hyperlinks"].append({"coordinate": cell.coordinate, **item["hyperlink"]})
                if cell.comment:
                    item["comment"] = {"author": cell.comment.author, "text": cell.comment.text}; sheet["comments"].append({"coordinate": cell.coordinate, **item["comment"]})
                sheet["cells"].append(item)
            for i, chart in enumerate(ws._charts):
                sheet["charts"].append(_chart_record(chart, i))
                result.embedded_objects.append({"kind": "chart", "sheet": ws.title, "index": i, "raw_data": sheet["charts"][-1], "visual_required": True})
            for i, image_obj in enumerate(ws._images):
                try: blob = image_obj._data()
                except Exception: blob = b""
                image_info = {"kind": "image", "sheet": ws.title, "index": i, "sha256": sha256_bytes(blob), "size": len(blob), "anchor": _anchor(image_obj.anchor)}
                sheet["images"].append(image_info); result.embedded_objects.append(image_info)
            sheets.append(sheet)
        names = []
        with zipfile.ZipFile(inp.path) as z:
            names = z.namelist()
            raw_drawings = [{"path": n, "sha256": sha256_bytes(z.read(n)), "size": z.getinfo(n).file_size} for n in names if n.startswith("xl/drawings/") and n.endswith(".xml")]
            chart_parts = [{"path": n, "sha256": sha256_bytes(z.read(n)), "size": z.getinfo(n).file_size} for n in names if n.startswith("xl/charts/") and n.endswith(".xml")]
        result.structured_semantics = {"sheet_count": len(sheets), "sheets": sheets, "named_ranges": [], "drawing_parts": raw_drawings, "chart_parts": chart_parts}
        try:
            defined = wb.defined_names.values() if hasattr(wb.defined_names, "values") else []
            result.structured_semantics["named_ranges"] = [{"name": x.name, "attr_text": x.attr_text, "localSheetId": x.localSheetId} for x in defined]
        except Exception:
            result.warnings.append("named range API unavailable")
        for sheet in sheets:
            _sheet_png(result, sheet, Path(output_dir) / "sheets" / f"sheet-{sheet['order']:04d}.png")
        formula_count = sum(1 for s in sheets for c in s["cells"] if c.get("formula"))
        chart_objects = [x for x in result.embedded_objects if x.get("kind") == "chart"]
        result.fidelity_flags = {"raw_data_fidelity": True, "formula_fidelity": True, "display_values_parser_derived": True, "sheet_order_and_hidden_preserved": True, "merged_ranges_preserved": True, "dimensions_preserved": True, "styles_accounted": True, "charts_accounted": all(x.get("raw_data") for x in chart_objects), "chart_visual_evidence": all(any(v.get("sheet") == x.get("sheet") for v in result.visual_pages) for x in chart_objects), "images_accounted": all(x.get("sha256") is not None for x in result.embedded_objects if x.get("kind") == "image"), "visual_evidence_available": bool(result.visual_pages), "formula_count": formula_count, "chart_count": len(chart_objects), "image_count": sum(1 for x in result.embedded_objects if x.get("kind") == "image"), "failure_explicit": True}
        result.warnings.append("sheet PNGs are deterministic structural visual proxies; native spreadsheet pagination is not required for data fidelity")
    except Exception as exc:
        return _read_xlsx_ooxml(inp, output_dir, parse_warning=str(exc))
    return result.finalize()


def read_xlsx(inp: ReaderInput, output_dir: str | Path) -> ReaderResult:
    return _read_xlsx_openpyxl(inp, output_dir)
