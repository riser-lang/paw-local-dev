"""Build standard exam corpus verbatim from tests/pending/exam/corpus.json.

Owned by executor in C-10. Generates docx, pptx, xlsx, md exactly as described in corpus.json.
Also provides semantic check evaluation helper for standard exam questions.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

DEFAULT_CORPUS_JSON = Path(__file__).resolve().parent.parent / "tests" / "pending" / "exam" / "corpus.json"
DEFAULT_QUESTIONS_JSON = Path(__file__).resolve().parent.parent / "tests" / "pending" / "exam" / "questions.json"


def load_corpus_spec(corpus_json_path: Path | None = None) -> list[dict]:
    path = corpus_json_path or DEFAULT_CORPUS_JSON
    spec = json.loads(path.read_text(encoding="utf-8"))
    return spec.get("documents", [])


def build_corpus(corpus_json_path: Path | None = None, output_dir: Path | None = None) -> list[Path]:
    import docx
    import openpyxl
    from pptx import Presentation

    if output_dir is None:
        output_dir = Path("tests/fixtures/exam")
    output_dir.mkdir(parents=True, exist_ok=True)

    documents = load_corpus_spec(corpus_json_path)
    built_paths: list[Path] = []

    for item in sorted(documents, key=lambda d: d.get("step", 0)):
        filename = item["file"]
        ftype = item["type"]
        target = output_dir / filename

        if ftype == "docx":
            doc = docx.Document()
            for p in item.get("paragraphs", []):
                doc.add_paragraph(p)
            doc.save(str(target))

        elif ftype == "pptx":
            prs = Presentation()
            for s_info in item.get("slides", []):
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = s_info.get("title", "")
                tf = slide.placeholders[1].text_frame
                body_lines = s_info.get("body", [])
                for i, text in enumerate(body_lines):
                    if i == 0:
                        tf.text = text
                    else:
                        p = tf.add_paragraph()
                        p.text = text
            prs.save(str(target))

        elif ftype == "xlsx":
            wb = openpyxl.Workbook()
            first = True
            for s_info in item.get("sheets", []):
                if first:
                    ws = wb.active
                    ws.title = s_info.get("name", "Sheet1")
                    first = False
                else:
                    ws = wb.create_sheet(title=s_info.get("name", "Sheet"))
                for row in s_info.get("rows", []):
                    ws.append(row)
            wb.save(str(target))

        elif ftype == "md":
            text = "\n\n".join(item.get("paragraphs", []))
            target.write_text(text, encoding="utf-8")

        else:
            raise ValueError(f"Unsupported corpus file type: {ftype} for {filename}")

        built_paths.append(target)

    return built_paths


def evaluate_semantics(
    answer: str,
    sources: list[str | dict],
    rules: dict,
    knowledge_gaps: list[dict] | None = None
) -> tuple[bool, str]:
    """Check assistant answer and sources against check_semantics rules."""
    src_names: list[str] = []
    for s in sources:
        if isinstance(s, str):
            src_names.append(Path(s).name)
        elif isinstance(s, dict):
            name = s.get("title") or s.get("file") or s.get("name") or s.get("filename") or ""
            src_names.append(Path(name).name)

    # 1. must_include
    if "must_include" in rules:
        for item in rules["must_include"]:
            if item not in answer:
                return False, f"must_include missing: '{item}'"

    # 2. must_include_any
    if "must_include_any" in rules:
        if not any(item in answer for item in rules["must_include_any"]):
            return False, f"must_include_any none matched: {rules['must_include_any']}"

    # 3. must_not_include
    if "must_not_include" in rules:
        for item in rules["must_not_include"]:
            if item in answer:
                return False, f"must_not_include forbidden item present: '{item}'"

    # 4. cite_file
    if "cite_file" in rules:
        target = Path(rules["cite_file"]).name
        if not any(target in s for s in src_names):
            return False, f"cite_file missing: '{target}' in sources {src_names}"

    # 5. must_not_cite_file
    if "must_not_cite_file" in rules:
        target = Path(rules["must_not_cite_file"]).name
        if any(target in s for s in src_names):
            return False, f"must_not_cite_file present: '{target}' in sources {src_names}"

    # 6. conflict_notice
    if "conflict_notice" in rules:
        vals = rules["conflict_notice"]
        if not all(v in answer for v in vals):
            return False, f"conflict_notice missing values {vals}"
        if not any(w in answer for w in ["冲突", "不一致", "矛盾"]):
            return False, "conflict_notice missing conflict keyword (冲突/不一致/矛盾)"

    # 7. suggestion_label
    if rules.get("suggestion_label"):
        if "建议" not in answer:
            return False, "suggestion_label missing '建议'"
        if "ANSWER_NOT_VERIFIED" in answer or "无法回答" in answer:
            return False, "suggestion_label refused answer"

    # 8. insufficient
    if rules.get("insufficient"):
        insuf_kw = ["依据不足", "没有找到", "资料中没有", "未提及", "未找到", "暂无"]
        if not any(k in answer for k in insuf_kw):
            return False, "insufficient keyword missing"
        if knowledge_gaps is not None:
            q = rules.get("q", "")
            if q and not any(q in g.get("question", "") or g.get("question", "") in q for g in knowledge_gaps):
                return False, f"insufficient question not recorded in knowledge_gaps: '{q}'"

    return True, "ok"


def main():
    parser = argparse.ArgumentParser(description="Build PAW standard exam corpus files")
    parser.add_argument(
        "--corpus-json",
        type=Path,
        default=DEFAULT_CORPUS_JSON,
        help="Path to corpus.json",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("tests/fixtures/exam"),
        help="Target directory to write corpus documents",
    )
    args = parser.parse_args()

    paths = build_corpus(args.corpus_json, args.output_dir)
    print(f"Built {len(paths)} exam documents in {args.output_dir}:")
    for p in paths:
        print(f"  - {p.name} ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
