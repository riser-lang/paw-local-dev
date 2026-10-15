"""PAW Desktop Local Edition - Complete Competition Demo Verification Script.

Demonstrates:
  File Import -> Reader Fidelity -> Evidence -> AI Candidate -> User Review -> Committer -> Formal State Update
"""
import json
import os
import sys
import time
from pathlib import Path

# Setup paths
root_path = Path(__file__).resolve().parent.parent
backend_path = root_path / "backend"
sys.path.insert(0, str(root_path))
sys.path.insert(0, str(backend_path))

import docx
from fastapi.testclient import TestClient
from main import app
from backend.core.config import DB_PATH, EVIDENCE_DIR, SOURCES_DIR
from backend.watcher.observer import watcher
from backend.governance.service import get_db

client = TestClient(app)


def print_step(num: int, title: str):
    print("\n" + "=" * 70)
    print(f"  STEP {num}: {title}")
    print("=" * 70)


def run_demo():
    print("""
    ===================================================================
    *           PAW Desktop Local Edition 参赛版本全链路演示          *
    *               Evidence First + Formal Governance                *
    ===================================================================
    """)

    # -------------------------------------------------------------
    # Step 1: Create a realistic business DOCX with structured elements
    # -------------------------------------------------------------
    print_step(1, "生成业务源文件并存入监控目录 (Drop Document into Sources)")
    doc_path = SOURCES_DIR / "PAW_2026_Q4_战略发展规划.docx"
    
    doc = docx.Document()
    doc.add_heading("PAW 2026 Q4 战略发展规划与产品落地方案", level=1)
    doc.add_paragraph("PAW 是基于证据保真 (Evidence-First) 与确定性治理的个人智能中枢。")
    doc.add_paragraph("核心决策：全面推进 Windows 本地参赛版本收敛，确保 Agent 永远不直接篡改正式状态。")
    
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "战略指标项"
    table.cell(0, 1).text = "目标预期"
    table.cell(1, 0).text = "正式知识点储备"
    table.cell(1, 1).text = "800+"
    table.cell(2, 0).text = "治理绕过违规次数 (formal_state_bypass)"
    table.cell(2, 1).text = "严格为 0"

    doc.add_paragraph("关键待办行动项：TODO: 组织决赛答辩模拟，现场演示高保真逆向与审核提交闭环。")
    doc.save(doc_path)
    print(f" -> 已生成文件: {doc_path}")
    print(f" -> 文件大小: {doc_path.stat().st_size} 字节")

    # -------------------------------------------------------------
    # Step 2: Source Watcher & Reader Fidelity Parsing
    # -------------------------------------------------------------
    print_step(2, "本地监听器自动感知 & Reader 高保真逆向提取 (Reader Fidelity)")
    from backend.reader.contract import artifact_sha256
    file_hash = artifact_sha256(doc_path)
    print(f" -> 原始文档 SHA256: {file_hash}")

    candidate_ids = watcher.process_file(doc_path, file_hash, event_type="competition_demo")
    print(f" -> Reader 提取成功! 生成证据包路径: {EVIDENCE_DIR}")
    print(f" -> AI 提议引擎 (Provider Adapter) 生成 Candidate 数量: {len(candidate_ids)}")
    for cid in candidate_ids:
        print(f"    * Candidate ID: {cid}")

    # -------------------------------------------------------------
    # Step 3: Inspect Evidence & Locators
    # -------------------------------------------------------------
    print_step(3, "审查证据包结构与定位符 (Inspect Evidence & Locators)")
    evidence_dir = list(EVIDENCE_DIR.glob(f"*{file_hash[:8]}"))
    if evidence_dir and (evidence_dir[0] / "reader_result.json").is_file():
        evidence_content = json.loads((evidence_dir[0] / "reader_result.json").read_text(encoding="utf-8"))
        print(f" -> 证据架构版本: {evidence_content.get('evidence_schema_version')}")
        print(f" -> 内容指纹: {evidence_content.get('content_fingerprint')}")
        print(f" -> 提取文本块数量: {len(evidence_content.get('text_blocks', []))}")
        print(f" -> 提取表格数量: {len(evidence_content.get('tables', []))}")
        for b in evidence_content.get("text_blocks", [])[:3]:
            print(f"    - [块 {b.get('index')}] {b.get('text')}")

    # -------------------------------------------------------------
    # Step 4: Web Workbench Review Center (/review)
    # -------------------------------------------------------------
    print_step(4, "前端审核中心获取待审提案 (Fetch Pending Candidates via API)")
    r_rev = client.get("/v1/web/brain/review/candidates?status=pending&limit=10")
    items = r_rev.json().get("items", [])
    demo_candidates = [item for item in items if item.get("candidate_id") in candidate_ids]
    print(f" -> 审核中心成功加载 {len(demo_candidates)} 条待审提案：")
    for dc in demo_candidates:
        pres = dc.get("presentation", {})
        print(f"    * [{pres.get('category_label')}] {pres.get('display_title')}")
        print(f"      - 影响实体: {pres.get('affected_entities')}")
        print(f"      - 证据支撑: {pres.get('evidence_items')}")

    # -------------------------------------------------------------
    # Step 5: User Approval & Deterministic Commit
    # -------------------------------------------------------------
    print_step(5, "用户执行审批 -> Committer 确定性提交 (Review -> Commit)")
    commit_receipts = []
    for target_candidate in demo_candidates:
        target_id = target_candidate["candidate_id"]
        pres = target_candidate.get("presentation", {})
        print(f" -> 用户核准提案: [{pres.get('category_label')}] {pres.get('display_title')}")
        r_accept = client.post(f"/v1/web/brain/review/candidates/{target_id}/accept", json={"review_notes": "评委现场核准写入"})
        assert r_accept.status_code == 200, f"Accept failed: {r_accept.text}"
        receipt = r_accept.json().get("result", {})
        commit_receipts.append((target_id, receipt))
        print(f"    * Committer 提交成功: Target={receipt.get('target_type')} (ID: {receipt.get('target_id')}) | CommitID: {receipt.get('commit_id')}")

    # -------------------------------------------------------------
    # Step 6: Verify Formal State in SQLite
    # -------------------------------------------------------------
    print_step(6, "正式业务状态校验 (Formal State 唯一性核验)")
    conn = get_db()
    try:
        for target_id, receipt in commit_receipts:
            cr = conn.execute("SELECT * FROM commit_record WHERE id = ?", (receipt.get("commit_id"),)).fetchone()
            assert cr is not None, "CommitRecord missing in formal state!"
            print(f" -> [DB] CommitRecord 持久化确认: result={cr['result']} (Change: {cr['change_id']})")

            cand = conn.execute("SELECT status, applied_entity_id FROM candidate WHERE id = ?", (target_id,)).fetchone()
            assert cand["status"] == "applied", "Candidate status not applied!"
            print(f" -> [DB] Candidate {target_id[:8]}... 状态已为 applied (已绑定实体: {cand['applied_entity_id']})")

        # Verify formal_state_bypass = 0
        total_commits = conn.execute("SELECT count(*) FROM commit_record WHERE result = 'success'").fetchone()[0]
        total_approved = conn.execute("SELECT count(*) FROM approved_change WHERE status = 'committed'").fetchone()[0]
        orphan_commits = conn.execute("SELECT count(*) FROM commit_record cr LEFT JOIN approved_change ac ON cr.change_id = ac.change_id WHERE ac.change_id IS NULL").fetchone()[0]
        
        print("\n" + "-" * 50)
        print("  PAW 严肃治理指标核验 (Governance Metrics):")
        print(f"   * 全库成功提交数: {total_commits}")
        print(f"   * 全库已批准变更数: {total_approved}")
        print(f"   * 孤儿非法写入数: {orphan_commits}")
        print(f"   * formal_state_bypass = {orphan_commits} (严格为 0)")
        print("-" * 50)
        assert orphan_commits == 0, "Security violation: found orphaned commits!"
    finally:
        conn.close()

    # -------------------------------------------------------------
    # Step 7: Evolution Engine Verification (Health, Graph, Memory)
    # -------------------------------------------------------------
    print_step(7, "Evolution Engine 自我演进与健康核验 (Health & Context Graph)")
    r_health = client.get("/v1/local/evolution/health")
    assert r_health.status_code == 200
    h_data = r_health.json()
    print(f" -> 知识资产整体健康度: {h_data['overall_score']} 分 ({h_data['overall_grade']})")
    print(f" -> 评估知识页面: {h_data['total_pages_evaluated']} 条 | 分布: {h_data['distribution']}")

    r_graph = client.get("/v1/local/evolution/graph?query=PAW")
    assert r_graph.status_code == 200
    g_data = r_graph.json()
    print(f" -> 上下文图谱 (Context Graph): 命中实体 {g_data['total_nodes']} 个, 跨实体关联 {g_data['total_edges']} 条")

    print("\n[SUCCESS] 比赛版本全链路验收演示 (含 Evolution Engine) 圆满成功！")


if __name__ == "__main__":
    run_demo()

