"""C-04 installed-app cases: real inputs, governance readback, real reader coverage."""
import hashlib
import json
import os
from pathlib import Path
from tests.run_pm_tests import (api,BASE,running,project,read_rows,dossier,progress,
    INSTALL_DIR,clean_and_record,silent_install,configure_mock_settings,poll_candidates_from_scanner)
from tests.run_pm_versions_tests import import_doc
from tests.run_e2e_tests import timed_answer


def current(pid):return progress(dossier(pid))['current']
def assessments(pid,sid):return [a for a in api('GET','/v1/pm/projects/'+pid+'/assessments')['items'] if a['source_version_id']==sid]
def conflicts(pid):return api('GET','/v1/pm/projects/'+pid+'/conflicts')['items']


def seed(case,text='原型已完成'):
    pid=project()
    sid,_=import_doc(case,'方案_v1.docx','冲刺项目。\n进度与里程碑：'+text+'。')
    initial=current(pid)
    if not initial or text not in initial['text']:raise RuntimeError('首份资料未通过正常语义/治理链成为结论')
    return pid,sid,initial


def make_conflict(case):
    pid,sid,initial=seed(case,'里程碑状态为已完成')
    new,_=import_doc(case,'现场记录.docx','冲刺项目。\n进度与里程碑：里程碑状态为尚未完成。')
    rows=conflicts(pid)
    if len(rows)!=1:raise RuntimeError('真实资料未产生唯一待裁决冲突')
    return pid,sid,new,initial,rows[0]


def capture_conflicts(c,pid,width):
    from playwright.sync_api import sync_playwright
    url=BASE+'/app/pm/projects/'+pid
    output=Path('screenshots')/('screen_'+c.id.replace('-','_')+'_'+str(width)+'.png')
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=False)
        try:
            page=browser.new_page(viewport={'width':width,'height':900})
            page.goto(url,wait_until='networkidle')
            page.get_by_role('heading',name='冲突清单',exact=True).wait_for()
            text=page.inner_text('body')
            c.check('conflict_two_sides_rendered_'+str(width),'另一份有效资料' in text and '里程碑状态为尚未完成' in text and '里程碑状态为已完成' in text,True)
            if width==390:c.check('conflicts_mobile_no_overflow',page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'),True)
            page.screenshot(path=str(output),full_page=True)
        finally:browser.close()
    c.screenshot(output,url)


def decide_in_ui(pid):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=False)
        try:
            page=browser.new_page(viewport={'width':1440,'height':900})
            page.goto(BASE+'/app/pm/projects/'+pid,wait_until='networkidle')
            page.get_by_role('button',name='采纳新资料',exact=True).click()
            page.get_by_role('heading',name='进度与里程碑 · ✓ 已裁决',exact=True).wait_for()
            return page.inner_text('body')
        finally:browser.close()


def coverage_inputs(case):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pypdf import PdfReader,PdfWriter
    directory=Path(os.environ['TEMP'])/('PAW_PM_COVERAGE_'+case);directory.mkdir(parents=True,exist_ok=True)
    doc=Document();doc.add_paragraph('冲刺项目。进度与里程碑：原型已完成。');doc.add_table(rows=1,cols=1).cell(0,0).text='结构表内容';doc.save(directory/'说明.docx')
    book=Workbook();book.active['A1']='冲刺项目';book.active['B1']='已交付';book.save(directory/'数据.xlsx')
    slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='冲刺项目';slide.placeholders[1].text='可读文字';slides.slides.add_slide(slides.slide_layouts[6]);slides.save(directory/'展示.pptx')
    writer=PdfWriter();writer.add_page(PdfReader('tests/fixtures/sample_report.pdf').pages[0]);writer.add_blank_page(width=612,height=792)
    with (directory/'阅读.pdf').open('wb') as output:writer.write(output)
    hashes={path.suffix.upper().lstrip('.'):hashlib.sha256(path.read_bytes()).hexdigest() for path in directory.iterdir()}
    configure_mock_settings(directory)
    api('POST','/v1/local/scanner/start',{'directory':str(directory),'force_restart':True})
    poll_candidates_from_scanner()
    return hashes


def run_c04(report,installer):
    for case in ['PM-B3-01','PM-B3-02','PM-B3-03','PM-B3-04','PM-B3-05','PM-B3-06','PM-B4-01','PM-B4-02']:
        with report.case(case,'栏目判定与覆盖率：'+case) as c:
            c.fresh_install(clean_and_record,lambda:silent_install(installer)[2],INSTALL_DIR)
            with running():
                if case=='PM-B3-01':
                    pid,first,initial=seed(case)
                    count=read_rows('SELECT COUNT(*) AS n FROM pm_section_revision')[0]['n']
                    second,_=import_doc(case,'现场记录.docx','冲刺项目。\n进度与里程碑：原型已完成。')
                    c.check('consistency_compares_distinct_sources',second!=first,True)
                    decision=assessments(pid,second)[0]
                    c.check('relation_consistent',decision['relation'],'consistent')
                    c.check('consistent_kept_current',current(pid)['id'],initial['id'])
                    c.check('consistent_no_new_conclusion',read_rows('SELECT COUNT(*) AS n FROM pm_section_revision')[0]['n'],count)
                    c.check('consistent_no_user_decision',decision['status'],'consistent')
                elif case=='PM-B3-02':
                    pid,first,initial=seed(case)
                    second,_=import_doc(case,'现场记录.docx','冲刺项目。\n进度与里程碑：联调已启动。')
                    decision=assessments(pid,second)[0];updated=current(pid)
                    c.check('relation_supplement',decision['relation'],'supplement')
                    c.check('supplement_retains_and_adds','原型已完成' in updated['text'] and '联调已启动' in updated['text'],True)
                    c.check('both_sources_preserved',{r['source_version_id'] for r in updated['source_refs']},{first,second})
                    for source in updated['source_refs']:
                        c.check('source_opens_'+source['source_version_id'][:8],source['text'].strip() in api('GET',source['source_url'])['body'],True)
                    receipt=read_rows('SELECT r.review_channel,r.review_notes,a.status FROM approved_change a JOIN review_decision r ON r.id=a.approval_ref WHERE a.change_id=?',(decision['change_id'],))[0]
                    c.check('supplement_committed_by_R3',receipt['status']=='committed' and receipt['review_channel']=='standing_user_rule' and 'R3' in receipt['review_notes'],True)
                    api('POST','/v1/pm/changes/'+decision['change_id']+'/undo',{})
                    c.check('supplement_undo_restores',current(pid)['id'],initial['id'])
                elif case=='PM-B3-03':
                    pid,first,initial=seed(case)
                    second,_=import_doc(case,'方案_v2.docx','冲刺项目。\n进度与里程碑：原型已完成，联调已启动。')
                    decision=assessments(pid,second)[0]
                    c.check('same_family_relation_replace',decision['relation'],'replace')
                    c.check('replacement_auto_applied',decision['status'],'applied')
                    c.check('replacement_uses_new_source',current(pid)['source_version_id'],second)
                    c.check('replacement_actual_new_content','联调已启动' in current(pid)['text'],True)
                    c.check('replacement_old_in_history',initial['id'] in [r['id'] for r in progress(dossier(pid))['history']],True)
                elif case in ('PM-B3-04','PM-B3-06'):
                    pid,first,second,initial,conflict=make_conflict(case)
                    c.check('conflict_open',conflict['status'],'open')
                    c.check('conflict_keeps_current',current(pid)['id'],initial['id'])
                    c.check('conflict_sources_are_distinct',{conflict['left_source_id'],conflict['right_source_id']},{first,second})
                    if case=='PM-B3-04':
                        session=api('POST','/v1/web/agy/chat',{'mode':'chat','project_id':pid})['session_id']
                        result=timed_answer(session,'冲刺项目里程碑状态现在如何？',case)
                        c.check('assistant_discloses_unresolved_conflict','存在冲突' in result['response'] and '待用户裁决' in result['response'],True)
                        c.check('assistant_conflict_metadata',result['conflicts'][0]['id'],conflict['id'])
                        capture_conflicts(c,pid,1440);capture_conflicts(c,pid,390)
                    else:
                        text=decide_in_ui(pid)
                        resolved=conflicts(pid)[0]
                        c.check('conflict_resolved',resolved['status'],'resolved')
                        c.check('decision_really_updates_current',current(pid)['text'],'里程碑状态为尚未完成')
                        receipt=read_rows('SELECT r.approval_source,a.status FROM approved_change a JOIN review_decision r ON r.id=a.approval_ref WHERE a.change_id=?',(resolved['resolved_by'],))[0]
                        c.check('decision_governance',receipt,{'approval_source':'direct_user','status':'committed'})
                        c.check('decision_shown_in_ui','✓ 已裁决' in text,True)
                        api('POST','/v1/pm/changes/'+resolved['resolved_by']+'/undo',{})
                        c.check('resolution_undo_restores_current',current(pid)['id'],initial['id'])
                        c.check('resolution_undo_reopens_conflict',conflicts(pid)[0]['status'],'open')
                elif case=='PM-B3-05':
                    pid=project();before=dossier(pid)['project']
                    for index,text in enumerate(['新项目：新建配套项目','KPI 目标值改为 20','政策调整方案','绩效规则变更','人员选择方案','赛制安排'],1):
                        sid,_=import_doc(case,'事项_v'+str(index)+'.docx','冲刺项目。\n进度与里程碑：'+text+'。')
                        c.check('user_exclusive_requires_decision_'+str(index),assessments(pid,sid)[0]['status'],'needs_user')
                    c.check('user_exclusive_no_formal_conclusion',read_rows('SELECT COUNT(*) AS n FROM pm_section_revision')[0]['n'],0)
                    c.check('user_exclusive_no_new_project',read_rows('SELECT COUNT(*) AS n FROM project')[0]['n'],1)
                    c.check('user_exclusive_project_unchanged',dossier(pid)['project'],before)
                elif case=='PM-B4-01':
                    project();hashes=coverage_inputs(case)
                    rows=api('GET','/v1/pm/coverage')['items'];by_id={r['source_version_id']:r for r in rows}
                    c.check('all_four_sources_measured',set(by_id),set(hashes.values()))
                    for kind,sid in hashes.items():
                        row=by_id[sid]
                        c.check('reader_type_'+kind,row['reader_type'],kind)
                        c.check('actual_units_'+kind,bool(row['units']) and row['total_count']==len(row['units']),True)
                    c.check('word_unconfirmed_pagination_marked',any('分页未确认' in s for s in by_id[hashes['DOCX']]['limitations']),True)
                    c.check('spreadsheet_table_grid_recorded',any(u['kind']=='table' for u in by_id[hashes['XLSX']]['units']),True)
                    for kind in ('PDF','PPTX'):
                        unread=[u for u in by_id[hashes[kind]]['units'] if not u['read']]
                        c.check('unread_unit_explicit_'+kind,len(unread)==1 and bool(unread[0]['reason']),True)
                        c.check('coverage_ratio_'+kind,by_id[hashes[kind]]['ratio'],.5)
                else:
                    pid,first,initial=seed(case)
                    c.check('no_trigger_zero_rereads',api('POST','/v1/pm/reread/run',{'limit':5})['processed'],0)
                    for version in range(2,8):import_doc(case,'方案_v'+str(version)+'.docx','冲刺项目。\n进度与里程碑：阶段 '+str(version)+' 已启动。')
                    jobs=api('GET','/v1/pm/reread')['items']
                    c.check('only_actual_triggers_queued',len(jobs)==6 and all(j['reason']=='new_version' for j in jobs),True)
                    first_round=api('POST','/v1/pm/reread/run',{'limit':5})
                    c.check('reread_round_capped_at_five',first_round['processed'],5)
                    c.check('real_rereads_completed',all(j['status']=='completed' for j in first_round['items']),True)
                    second_round=api('POST','/v1/pm/reread/run',{'limit':5})
                    c.check('reread_remaining_once',second_round['processed'],1)
                    c.check('completed_triggers_not_repeated',api('POST','/v1/pm/reread/run',{'limit':5})['processed'],0)
                    final=api('GET','/v1/pm/reread')['items']
                    c.check('reread_real_reader_and_coverage',all(json.loads(j['result_json']).get('reader_type')=='DOCX' and json.loads(j['result_json']).get('coverage',{}).get('read_count') for j in final),True)
